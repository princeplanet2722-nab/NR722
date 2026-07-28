#!/usr/bin/env python3
"""Build DOCX + PDF deliverables from the AI Verse markdown scripts.

Usage:  python3 aiverse/tools/build_docs.py
Requires: python-docx, reportlab
"""
import os
import re
import html

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer,
                                HRFlowable, Table, TableStyle)

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
SRC = os.path.join(ROOT, "aiverse", "script")
OUT = os.path.join(ROOT, "aiverse", "downloads")

FILES = [
    ("5-secret-ai-tools-2026-SCRIPT.md", "AI-Verse-5-Secret-AI-Tools-2026-SCRIPT"),
    ("VIDEO-PACKAGE.md", "AI-Verse-5-Secret-AI-Tools-2026-UPLOAD-PACKAGE"),
]

ACCENT = RGBColor(0x00, 0xB4, 0xD8)
CUE = RGBColor(0xC0, 0x39, 0x2B)


def parse(md_text):
    """Yield (kind, payload) blocks from a restricted markdown subset."""
    lines = md_text.split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("```"):
            buf = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            yield ("code", "\n".join(buf))
            continue
        if line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not set("".join(cells)) <= set("-: "):
                    rows.append(cells)
                i += 1
            yield ("table", rows)
            continue
        if not line.strip():
            i += 1
            continue
        if re.match(r"^(-{3,}|\*{3,})$", line.strip()):
            yield ("rule", "")
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            yield ("h%d" % min(level, 4), line.lstrip("#").strip())
        elif line.startswith(">"):
            yield ("quote", line.lstrip("> ").strip())
        elif re.match(r"^\s*[-*]\s+", line):
            yield ("bullet", re.sub(r"^\s*[-*]\s+", "", line))
        else:
            yield ("p", line.strip())
        i += 1


def clean(text):
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    text = re.sub(r"\[(.+?)\]\((.+?)\)", r"\1", text)
    return text


# ---------------------------------------------------------------- DOCX
def build_docx(md_text, out_path, title):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.7)
        s.left_margin = s.right_margin = Inches(0.8)

    for kind, payload in parse(md_text):
        if kind == "rule":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            continue
        if kind == "code":
            for ln in payload.split("\n"):
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.3)
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(ln)
                r.font.name = "Consolas"
                r.font.size = Pt(9.5)
            doc.add_paragraph()
            continue
        if kind == "table":
            rows = payload
            t = doc.add_table(rows=len(rows), cols=len(rows[0]))
            t.style = "Table Grid"
            for ri, row in enumerate(rows):
                for ci, cell in enumerate(row[: len(rows[0])]):
                    cp = t.cell(ri, ci).paragraphs[0]
                    r = cp.add_run(clean(cell))
                    r.font.size = Pt(9.5)
                    if ri == 0:
                        r.bold = True
            doc.add_paragraph()
            continue

        text = clean(payload)
        if kind.startswith("h"):
            lvl = int(kind[1])
            p = doc.add_paragraph()
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt({1: 20, 2: 15, 3: 12.5, 4: 11.5}[lvl])
            if lvl <= 2:
                r.font.color.rgb = ACCENT
            p.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(5)
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(12)
        elif kind == "bullet":
            p = doc.add_paragraph(text, style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            # colour the [CUE] tags
            for part in re.split(r"(\[[^\]]+\])", text):
                if not part:
                    continue
                r = p.add_run(part)
                if part.startswith("["):
                    r.bold = True
                    r.font.color.rgb = CUE
    doc.save(out_path)


# ---------------------------------------------------------------- PDF
def build_pdf(md_text, out_path, title):
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName="Helvetica",
                          fontSize=10, leading=14, spaceAfter=5)
    quote = ParagraphStyle("quote", parent=body, fontName="Helvetica-Bold",
                           fontSize=11, leading=15, leftIndent=16, spaceAfter=8,
                           textColor=colors.HexColor("#111111"))
    code = ParagraphStyle("code", parent=body, fontName="Courier", fontSize=8.5,
                          leading=11, leftIndent=14, spaceAfter=1,
                          backColor=colors.HexColor("#F4F4F4"))
    heads = {
        1: ParagraphStyle("h1", parent=ss["Heading1"], fontSize=19, leading=23,
                          textColor=colors.HexColor("#0077B6"), spaceBefore=10, spaceAfter=6),
        2: ParagraphStyle("h2", parent=ss["Heading2"], fontSize=14, leading=18,
                          textColor=colors.HexColor("#0096C7"), spaceBefore=12, spaceAfter=5),
        3: ParagraphStyle("h3", parent=ss["Heading3"], fontSize=11.5, leading=15,
                          spaceBefore=9, spaceAfter=4),
        4: ParagraphStyle("h4", parent=ss["Heading4"], fontSize=10.5, leading=14,
                          spaceBefore=8, spaceAfter=3),
    }
    doc = SimpleDocTemplate(out_path, pagesize=LETTER, title=title,
                            author="AI Verse", leftMargin=0.75 * inch,
                            rightMargin=0.75 * inch, topMargin=0.7 * inch,
                            bottomMargin=0.7 * inch)
    flow = []

    def esc(t):
        t = html.escape(clean(t))
        return re.sub(r"(\[[^\]]+\])", r'<font color="#C0392B"><b>\1</b></font>', t)

    for kind, payload in parse(md_text):
        if kind == "rule":
            flow.append(Spacer(1, 4))
            flow.append(HRFlowable(width="100%", thickness=0.6,
                                   color=colors.HexColor("#CCCCCC")))
            flow.append(Spacer(1, 6))
        elif kind == "code":
            for ln in payload.split("\n"):
                flow.append(Paragraph(html.escape(ln) or "&nbsp;", code))
            flow.append(Spacer(1, 8))
        elif kind == "table":
            rows = payload
            data = [[Paragraph("<b>%s</b>" % html.escape(clean(c)) if ri == 0
                               else html.escape(clean(c)),
                               ParagraphStyle("c", parent=body, fontSize=8.5, leading=11))
                     for c in row] for ri, row in enumerate(rows)]
            t = Table(data, colWidths=[1.5 * inch, 2.2 * inch, 3.0 * inch][: len(rows[0])])
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8F6FB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            flow.append(t)
            flow.append(Spacer(1, 10))
        elif kind.startswith("h"):
            flow.append(Paragraph(esc(payload), heads[int(kind[1])]))
        elif kind == "quote":
            flow.append(Paragraph("&ldquo;%s&rdquo;" % esc(payload), quote))
        elif kind == "bullet":
            flow.append(Paragraph("&bull;&nbsp; %s" % esc(payload),
                                  ParagraphStyle("b", parent=body, leftIndent=12)))
        else:
            flow.append(Paragraph(esc(payload), body))
    doc.build(flow)


def main():
    os.makedirs(OUT, exist_ok=True)
    for src, stem in FILES:
        with open(os.path.join(SRC, src), encoding="utf-8") as f:
            md = f.read()
        title = stem.replace("-", " ")
        build_docx(md, os.path.join(OUT, stem + ".docx"), title)
        build_pdf(md, os.path.join(OUT, stem + ".pdf"), title)
        print("built:", stem + ".docx", "+", stem + ".pdf")


if __name__ == "__main__":
    main()
